"""
add_locations.py — append admin location names to a customer workbook from its
lat/long columns, using the same shapefiles as the dashboard (offline, no API).

    python add_locations.py "ACTIVE CUSTOMERS 2026 RB.xlsx" [--sheet Sheet1]

Adds COUNTY, SUB_COUNTY, WARD (Kenya_Wards) and DIVISION, LOCATION, SUBLOCATION,
SALES_TERRITORY (Boroughs) as new columns on the same sheet. Existing cells and
formatting are untouched; a .bak copy of the original is written first.
Points outside every polygon are left blank. Re-running replaces the columns.
"""

import argparse
import shutil
import warnings

import geopandas as gpd
import openpyxl
import pandas as pd

warnings.filterwarnings("ignore", message="Geometry is in a geographic CRS")

WARDS_SHP   = "Kenya_Wards/kenya_wards.shp"
BOROUGH_SHP = "Boroughs/boroughs and branches.shp"
NEW_COLS = ["COUNTY", "SUB_COUNTY", "WARD",
            "DIVISION", "LOCATION", "SUBLOCATION", "SALES_TERRITORY"]


def _find_col(columns, *names):
    lower = {str(c).strip().lower(): c for c in columns}
    for n in names:
        if n in lower:
            return lower[n]
    raise SystemExit(f"No column named any of {names} in {list(columns)}")


def _join(points, polys, cols):
    """Point-in-polygon join; first match per point, indexed like `points`."""
    j = gpd.sjoin(points, polys[cols + ["geometry"]], how="left", predicate="within")
    return j[~j.index.duplicated(keep="first")][cols]


def locate(df):
    lat = pd.to_numeric(df[_find_col(df.columns, "lat", "latitude")], errors="coerce")
    lon = pd.to_numeric(df[_find_col(df.columns, "long", "lon", "lng", "longitude")], errors="coerce")
    ok = lat.notna() & lon.notna()
    pts = gpd.GeoDataFrame(index=df.index[ok],
                           geometry=gpd.points_from_xy(lon[ok], lat[ok]), crs="EPSG:4326")

    wards = gpd.read_file(WARDS_SHP).to_crs("EPSG:4326")
    w = _join(pts, wards, ["county", "subcounty", "ward"])

    bor = gpd.read_file(BOROUGH_SHP).to_crs("EPSG:4326")
    b = _join(pts, bor, ["DIVNAME", "LOCNAME", "SLNAME", "Boroughs"])

    out = pd.DataFrame(index=df.index, columns=NEW_COLS, dtype=object)
    out.loc[w.index, ["COUNTY", "SUB_COUNTY", "WARD"]] = w.values
    out.loc[b.index, ["DIVISION", "LOCATION", "SUBLOCATION", "SALES_TERRITORY"]] = b.values
    for c in NEW_COLS:
        out[c] = out[c].where(out[c].isna(), out[c].astype(str).str.strip().str.title())
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--sheet", default=None, help="default: first sheet")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.workbook)
    ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
    df = pd.read_excel(a.workbook, sheet_name=ws.title)
    loc = locate(df)

    shutil.copy2(a.workbook, a.workbook + ".bak")

    # Reuse existing columns on re-run, else append after the last one.
    header = {str(c.value): c.column for c in ws[1] if c.value is not None}
    next_col = ws.max_column + 1
    for name in NEW_COLS:
        col = header.get(name)
        if col is None:
            col, next_col = next_col, next_col + 1
            ws.cell(row=1, column=col, value=name)
        for i, v in enumerate(loc[name].tolist(), start=2):   # data starts on row 2
            ws.cell(row=i, column=col, value=None if pd.isna(v) else v)
    wb.save(a.workbook)

    print(f"{a.workbook} [{ws.title}] — {len(df):,} rows")
    for c in NEW_COLS:
        print(f"  {c:<16} {loc[c].notna().sum():>6,} matched  {loc[c].isna().sum():>5,} blank")


if __name__ == "__main__":
    main()
